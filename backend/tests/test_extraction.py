import tempfile
from pathlib import Path
import pytest
import pypdf
import docx

from backend.app.ingestion.extractor import DocumentExtractor
from backend.app.ingestion.chunker import DocumentChunker


def test_txt_extraction_and_chunking():
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
        f.write("Line 1 of text.\n\nLine 2 with detailed explanation of architecture.")
        temp_path = Path(f.name)

    try:
        doc = DocumentExtractor.extract(temp_path, "doc_test_1", "sample.txt")
        assert doc.document_id == "doc_test_1"
        assert doc.filename == "sample.txt"
        assert doc.file_type == "txt"
        assert doc.total_pages >= 1

        chunker = DocumentChunker(chunk_size=50, chunk_overlap=10)
        chunks = chunker.chunk_document(doc)
        assert len(chunks) >= 1
        assert chunks[0].document_id == "doc_test_1"
        assert chunks[0].filename == "sample.txt"
        assert chunks[0].page_number == 1
        assert "sample.txt" in chunks[0].metadata["filename"]
    finally:
        temp_path.unlink(missing_ok=True)


def test_unsupported_file_extension():
    fake_path = Path("fake_document.xyz")
    with pytest.raises(ValueError, match="Unsupported file format"):
        DocumentExtractor.extract(fake_path, "doc_unsupported", "fake.xyz")


def test_empty_txt_raises_value_error():
    """Verify that empty text files raise ValueError."""
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
        f.write("   \n\n  \t  ")
        temp_path = Path(f.name)

    try:
        with pytest.raises(ValueError, match="No readable text could be extracted"):
            DocumentExtractor.extract(temp_path, "doc_empty", "empty.txt")
    finally:
        temp_path.unlink(missing_ok=True)


def test_scanned_or_blank_pdf_raises_value_error():
    """Verify that a PDF without any extractable text (e.g. image-only scan) raises ValueError."""
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with tempfile.NamedTemporaryFile("wb", suffix=".pdf", delete=False) as f:
        writer.write(f)
        temp_path = Path(f.name)

    try:
        with pytest.raises(ValueError, match="No readable text could be extracted"):
            DocumentExtractor.extract(temp_path, "doc_scanned", "scanned.pdf")
    finally:
        temp_path.unlink(missing_ok=True)


def test_corrupt_pdf_raises_value_error():
    """Verify that corrupted PDF files raise ValueError."""
    with tempfile.NamedTemporaryFile("wb", suffix=".pdf", delete=False) as f:
        f.write(b"%PDF-1.4\nCorrupt trailer invalid binary \x00\xff\xfe")
        temp_path = Path(f.name)

    try:
        with pytest.raises(ValueError, match="Corrupted or unreadable PDF file"):
            DocumentExtractor.extract(temp_path, "doc_corrupt_pdf", "corrupt.pdf")
    finally:
        temp_path.unlink(missing_ok=True)


def test_corrupt_docx_raises_value_error():
    """Verify that corrupted DOCX files raise ValueError."""
    with tempfile.NamedTemporaryFile("wb", suffix=".docx", delete=False) as f:
        f.write(b"PK\x03\x04not a valid zip package header")
        temp_path = Path(f.name)

    try:
        with pytest.raises(ValueError, match="Corrupted or unreadable DOCX file"):
            DocumentExtractor.extract(temp_path, "doc_corrupt_docx", "corrupt.docx")
    finally:
        temp_path.unlink(missing_ok=True)


def test_docx_table_content_extraction():
    """Verify that tables inside DOCX documents are extracted and included in chunks."""
    doc = docx.Document()
    doc.add_paragraph("Architecture Overview:")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Role"
    table.rows[1].cells[0].text = "Risk Engine"
    table.rows[1].cells[1].text = "Human Review Escalation"

    with tempfile.NamedTemporaryFile("wb", suffix=".docx", delete=False) as f:
        doc.save(f)
        temp_path = Path(f.name)

    try:
        extracted = DocumentExtractor.extract(temp_path, "doc_docx_tbl", "spec.docx")
        combined_text = "\n".join(p.text for p in extracted.pages)
        assert "Risk Engine" in combined_text
        assert "Human Review Escalation" in combined_text
        assert "Component | Role" in combined_text
    finally:
        temp_path.unlink(missing_ok=True)


def test_multipage_pdf_page_number_preservation():
    """Verify page number preservation across pages in PDF."""
    from reportlab.pdfgen import canvas

    with tempfile.NamedTemporaryFile("wb", suffix=".pdf", delete=False) as f:
        c = canvas.Canvas(f.name)
        c.drawString(100, 750, "Content for Page One: Feature Extraction")
        c.showPage()
        c.drawString(100, 750, "Content for Page Two: Synthetic Speech Threshold 0.75")
        c.showPage()
        c.drawString(100, 750, "Content for Page Three: Measured Accuracy is unstated")
        c.showPage()
        c.save()
        temp_path = Path(f.name)

    try:
        extracted = DocumentExtractor.extract(temp_path, "doc_multipage", "voice_shield.pdf")
        assert extracted.total_pages == 3
        assert len(extracted.pages) == 3
        assert extracted.pages[0].page_number == 1
        assert extracted.pages[1].page_number == 2
        assert extracted.pages[2].page_number == 3
        assert "Page One" in extracted.pages[0].text
        assert "0.75" in extracted.pages[1].text

        chunker = DocumentChunker()
        chunks = chunker.chunk_document(extracted)
        chunk_pages = [c.page_number for c in chunks]
        assert 1 in chunk_pages
        assert 2 in chunk_pages
        assert 3 in chunk_pages
    finally:
        temp_path.unlink(missing_ok=True)
