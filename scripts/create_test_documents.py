import os
from pathlib import Path
from docx import Document

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
except ImportError:
    canvas = None


def create_pdf(output_path: Path):
    if not canvas:
        raise RuntimeError("ReportLab is required to generate sample PDF.")

    c = canvas.Canvas(str(output_path), pagesize=letter)
    width, height = letter

    # Page 1
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 72, "Acoustic Signal Preprocessing & Windowing")
    c.setFont("Helvetica", 11)
    text_p1 = [
        "Audio signals in speech processing are recorded in the continuous-time domain and digitized via sampling.",
        "To process non-stationary speech, signals are segmented into short frames of 20 to 30 milliseconds.",
        "Framing assumes local stationarity. Window functions like Hamming or Hanning are applied to each frame",
        "to reduce spectral leakage at the boundaries before computing frequency transforms.",
        "Pre-emphasis filtering is commonly used to boost high-frequency components that naturally attenuate."
    ]
    y = height - 110
    for line in text_p1:
        c.drawString(72, y, line)
        y -= 20
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(72, 50, "Page 1 - Audio Signal Preprocessing")
    c.showPage()

    # Page 2
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 72, "Short-Time Fourier Transform (STFT)")
    c.setFont("Helvetica", 11)
    text_p2 = [
        "The Short-Time Fourier Transform (STFT) computes the Discrete Fourier Transform (DFT) of each windowed frame.",
        "This transforms the 1D time-domain signal into a 2D complex time-frequency spectrogram.",
        "Squaring the magnitude of the STFT yields the linear power spectrum, showing energy distribution across frequency bins.",
        "Linear frequency bins give equal resolution to high and low frequencies, which diverges from human hearing."
    ]
    y = height - 110
    for line in text_p2:
        c.drawString(72, y, line)
        y -= 20
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(72, 50, "Page 2 - Short-Time Fourier Transform")
    c.showPage()

    # Page 3
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 72, "The Mel Scale and Mel-Spectrogram")
    c.setFont("Helvetica", 11)
    text_p3 = [
        "Human auditory perception perceives pitch non-linearly, having high sensitivity below 1000 Hz and lower sensitivity above.",
        "The Mel scale maps physical frequency f in Hz to perceptual pitch: m = 2595 * log10(1 + f / 700).",
        "A Mel filterbank consists of triangular overlapping bandpass filters spaced linearly below 1000 Hz and logarithmically above.",
        "Multiplying the linear power spectrum by the Mel filterbank produces the Mel-spectrogram.",
        "Crucial definition: A Mel-spectrogram is an energy distribution across correlated mel-frequency channels.",
        "It retains energy across overlapping auditory filter channels and is an intermediate spectral representation."
    ]
    y = height - 110
    for line in text_p3:
        c.drawString(72, y, line)
        y -= 20
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(72, 50, "Page 3 - Mel-Scale and Mel-Spectrogram")
    c.showPage()

    # Page 4
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 72, "MFCCs: Mel-Frequency Cepstral Coefficients")
    c.setFont("Helvetica", 11)
    text_p4 = [
        "Mel-Frequency Cepstral Coefficients (MFCCs) are derived from the Mel-spectrogram through two critical operations:",
        "1. Logarithmic compression: Taking the logarithm of the Mel filterbank energies to mirror perceptual human loudness.",
        "2. Discrete Cosine Transform (DCT-II): Applying the DCT across the log-mel filterbank energies.",
        "",
        "The Discrete Cosine Transform (DCT) decorrelates the filterbank energies and compresses spectral envelope information",
        "into orthogonal cepstral coefficients (typically the first 12 to 13 coefficients are retained).",
        "",
        "CRITICAL TECHNICAL DISTINCTION:",
        "A Mel-spectrogram is an intermediate spectral energy representation across correlated mel frequency channels.",
        "In contrast, MFCCs are decorrelated, orthogonal cepstral coefficients produced by applying the Discrete Cosine Transform (DCT).",
        "Therefore, a Mel-spectrogram is not identical to MFCCs; the DCT and logarithmic compression distinguish MFCCs from Mel-spectrograms."
    ]
    y = height - 110
    for line in text_p4:
        c.drawString(72, y, line)
        y -= 20
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(72, 50, "Page 4 - Mel-Frequency Cepstral Coefficients (MFCCs)")
    c.showPage()

    # Page 5
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, height - 72, "Feature Normalization & Speech Recognition")
    c.setFont("Helvetica", 11)
    text_p5 = [
        "In Automatic Speech Recognition (ASR), MFCC features are supplemented with delta (velocity) and delta-delta (acceleration) coefficients.",
        "Cepstral Mean and Variance Normalization (CMVN) is applied across utterance frames to remove channel distortions and microphone variability.",
        "This yields robust acoustic features for downstream Hidden Markov Models or deep neural network acoustic models."
    ]
    y = height - 110
    for line in text_p5:
        c.drawString(72, y, line)
        y -= 20
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(72, 50, "Page 5 - Feature Normalization in ASR")
    c.showPage()

    c.save()
    print(f"Created PDF at: {output_path}")


def create_docx(output_path: Path):
    doc = Document()
    doc.add_heading("Distributed Systems Architecture & Consensus", level=0)

    # Section 1
    doc.add_heading("CAP Theorem and PACELC Model", level=1)
    doc.add_paragraph(
        "In distributed data stores, Eric Brewer's CAP theorem states that in the event of a network partition (P), "
        "a system must choose between Consistency (C) and Availability (A). Consistency guarantees that every read "
        "receives the most recent write or an error. Availability guarantees that every non-failing node returns a response."
    )
    doc.add_paragraph(
        "Daniel Abadi extended this trade-off with the PACELC theorem: If there is a Partition (P), trade off Availability (A) "
        "and Consistency (C); Else (E), trade off Latency (L) and Consistency (C). Systems like DynamoDB emphasize latency and availability, "
        "whereas Spanner emphasizes strict serializable consistency."
    )

    # Section 2
    doc.add_heading("Vector Clocks and Conflict Resolution", level=1)
    doc.add_paragraph(
        "In decentralized peer-to-peer topologies lacking centralized atomic clocks, causality between events is tracked using vector clocks. "
        "A vector clock is an array of logical timestamps maintained by each node. When comparing two vector clocks V1 and V2, if neither dominates "
        "the other, concurrent writes have occurred, requiring application-level conflict resolution (e.g. CRDTs or Last-Write-Wins)."
    )

    # Section 3
    doc.add_heading("Raft Consensus Algorithm and Leader Election", level=1)
    doc.add_paragraph(
        "The Raft consensus algorithm is designed for understandability while providing equivalent fault tolerance to Multi-Paxos. "
        "Raft decomposes consensus into three independent subproblems: Leader Election, Log Replication, and Safety."
    )
    doc.add_paragraph(
        "A Raft node exists in one of three states: Leader, Follower, or Candidate. Leaders send periodic heartbeats to maintain authority. "
        "If a follower receives no heartbeats within an election timeout (typically randomized between 150ms and 300ms), it transitions to "
        "Candidate state, increments its current term, votes for itself, and issues RequestVote RPCs to all peers. A candidate becomes leader "
        "upon receiving votes from a majority (quorum) of nodes."
    )

    doc.save(str(output_path))
    print(f"Created DOCX at: {output_path}")


if __name__ == "__main__":
    out_dir = Path("data/sample_docs")
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = out_dir / "audio_processing_and_speech.pdf"
    docx_path = out_dir / "distributed_systems_architecture.docx"
    create_pdf(pdf_path)
    create_docx(docx_path)

